# -*- coding: utf-8 -*-
"""
Verification script for Healthkathon_Kandidat_1.docx
Checks paragraph count, headings, tables, and alignment with Aegis-JKN specifications.
"""

import docx

def verify_docx(file_path):
    doc = docx.Document(file_path)
    print(f"Verifying: {file_path}")
    print(f"Total Paragraphs: {len(doc.paragraphs)}")
    print(f"Total Tables: {len(doc.tables)}")

    headings_1 = []
    headings_2 = []

    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt.startswith("BAB ") or txt in ["RINGKASAN EKSEKUTIF", "DAFTAR ISI", "DAFTAR TABEL & GAMBAR", "DAFTAR PUSTAKA & REFERENSI REGULASI"]:
            headings_1.append(txt)
        elif any(txt.startswith(f"{i}.") for i in range(1, 13)):
            headings_2.append(txt)

    print("\n--- BAB / MAJOR HEADINGS FOUND ---")
    for h in headings_1:
        print(f"  • {h}")

    print(f"\nTotal Subheadings (Level 2): {len(headings_2)}")
    print("\nSample Subheadings:")
    for h in headings_2[:10]:
        print(f"  - {h}")

    # Check key terms
    full_text = " ".join([p.text for p in doc.paragraphs])
    key_terms = [
        "Aegis-JKN",
        "Multi-Hop Graph Structure Learning",
        "MHGSL",
        "Topological Camouflage",
        "Tri-Tier",
        "OASIS",
        "SYND-01",
        "SYND-09",
        "Rp 507.370.000",
        "Rp 3.526.650.000",
        "UU PDP No. 27/2022",
        "HL7 FHIR R4",
        "Safe-Mode Fallback",
        "Next.js 16"
    ]

    print("\n--- KEY TERM VALIDATION ---")
    for term in key_terms:
        count = full_text.count(term)
        status = "OK" if count > 0 else "MISSING"
        print(f"  [{status}] '{term}': {count} occurrences")

if __name__ == "__main__":
    verify_docx("Healthkathon_Kandidat_1.docx")
