# -*- coding: utf-8 -*-
"""
Script Generator Dokumen Proposal Healthkathon 2026: Aegis-JKN
Menyelaraskan Healthkathon_Kandidat_1.docx dengan Codebase Aegis-JKN, Teori MHGSL,
dan Dokumentasi Strategis Notion.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_document():
    doc = docx.Document()

    # 1. Page Setup (A4, 4-3-3-3 cm margins)
    sec = doc.sections[0]
    sec.page_width = Pt(595.45)   # A4 Width (210 mm)
    sec.page_height = Pt(841.7)   # A4 Height (297 mm)
    sec.left_margin = Pt(113.05)  # 4 cm
    sec.top_margin = Pt(84.95)    # 3 cm
    sec.right_margin = Pt(84.95)  # 3 cm
    sec.bottom_margin = Pt(84.95) # 3 cm

    # Header & Footer
    header = sec.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run("HEALTHKATHON BPJS KESEHATAN 2026 — PROPOSAL INOVASI AEGIS-JKN")
    hrun.font.name = "Times New Roman"
    hrun.font.size = Pt(8.5)
    hrun.font.italic = True
    hrun.font.color.rgb = RGBColor(100, 110, 125)

    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    frun = fp.add_run("Aegis-JKN: Multi-Hop Graph Structure Learning for JKN Fraud Intelligence")
    frun.font.name = "Times New Roman"
    frun.font.size = Pt(8.5)
    frun.font.color.rgb = RGBColor(120, 130, 140)

    return doc

# XML Helpers
def set_cell_background(cell, fill_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=140, right=140):
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    cell._tc.get_or_add_tcPr().append(tcMar)

def set_cell_borders(cell, left="B0C4DE", top="B0C4DE", right="B0C4DE", bottom="B0C4DE", sz="4"):
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{top}"/>
            <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{left}"/>
            <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{bottom}"/>
            <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{right}"/>
        </w:tcBorders>
    ''')
    cell._tc.get_or_add_tcPr().append(borders)

def add_callout(doc, tag, title, body_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    cell.width = Pt(397.45) # ~7920 dxa (fit within text boundaries)

    set_cell_background(cell, "F0F4F8")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=180)

    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="36" w:space="0" w:color="1F4E79"/>
            <w:top w:val="nil"/>
            <w:bottom w:val="nil"/>
            <w:right w:val="nil"/>
        </w:tcBorders>
    ''')
    cell._tc.get_or_add_tcPr().append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    rtag = p.add_run(f"[{tag}] ")
    rtag.font.name = "Times New Roman"
    rtag.font.size = Pt(10.5)
    rtag.bold = True
    rtag.font.color.rgb = RGBColor(31, 78, 121)

    rtitle = p.add_run(f"{title}\n")
    rtitle.font.name = "Times New Roman"
    rtitle.font.size = Pt(10.5)
    rtitle.bold = True
    rtitle.font.color.rgb = RGBColor(20, 45, 75)

    rbody = p.add_run(body_text)
    rbody.font.name = "Times New Roman"
    rbody.font.size = Pt(10.0)
    rbody.italic = True
    rbody.font.color.rgb = RGBColor(30, 40, 50)

    # Spacing after callout
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(13.5)
    r.bold = True
    r.font.color.rgb = RGBColor(31, 78, 121)
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12.0)
    r.bold = True
    r.font.color.rgb = RGBColor(31, 78, 121)
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12.0)
    r.bold = True
    r.italic = True
    r.font.color.rgb = RGBColor(40, 50, 60)
    return p

def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    if bold_prefix:
        rb = p.add_run(bold_prefix)
        rb.font.name = "Times New Roman"
        rb.font.size = Pt(12.0)
        rb.bold = True

    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12.0)
    return p

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    if bold_prefix:
        rb = p.add_run(bold_prefix)
        rb.font.name = "Times New Roman"
        rb.font.size = Pt(12.0)
        rb.bold = True

    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12.0)
    return p

def add_table_data(doc, headers, data, col_widths=None):
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Format header row
    hdr_row = tbl.rows[0]
    hdr_row._tr.get_or_add_trPr().append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    hdr_row._tr.get_or_add_trPr().append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

    for c_idx, title in enumerate(headers):
        cell = hdr_row.cells[c_idx]
        set_cell_background(cell, "1F4E79")
        set_cell_margins(cell, top=120, bottom=120, left=120, right=120)
        set_cell_borders(cell, left="B0C4DE", top="1F4E79", right="B0C4DE", bottom="1F4E79", sz="6")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(title)
        r.font.name = "Times New Roman"
        r.font.size = Pt(9.5)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    # Format data rows
    for r_idx, row_values in enumerate(data):
        row = tbl.rows[r_idx + 1]
        row._tr.get_or_add_trPr().append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        bg_color = "F7F9FB" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_values):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            set_cell_borders(cell, left="D0DCE5", top="D0DCE5", right="D0DCE5", bottom="D0DCE5", sz="4")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(str(val))
            r.font.name = "Times New Roman"
            r.font.size = Pt(9.0)
            r.font.color.rgb = RGBColor(20, 30, 40)

    # Set column widths if provided
    if col_widths:
        for r in tbl.rows:
            for idx, w in enumerate(col_widths):
                r.cells[idx].width = Pt(w)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(2)
    sp.paragraph_format.space_after = Pt(4)
    return tbl

print("Helper functions initialized.")
